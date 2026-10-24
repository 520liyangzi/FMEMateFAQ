"""Behavior tests use small in-memory workbooks, never real customer files."""
import io
import unittest
import zipfile

import openpyxl

from fmemate.importer import DEFAULT_COLUMNS, preview_import

HEADERS = ["触发问题", "英文触发问题", "对象", "对象英文", "字段中文", "字段英文"]


def workbook_bytes(rows, merges=(), title="功能清单", prefix_sheet=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    if prefix_sheet:
        sheet.title = "填写说明"
        sheet.append(["请在功能清单工作表填写预置功能"])
        sheet = book.create_sheet(title)
    for row in rows:
        sheet.append(row)
    for merged in merges:
        sheet.merge_cells(merged)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class ImporterTests(unittest.TestCase):
    def test_vertical_merges_propagate_ancestry_and_multiple_objects(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "故障端口表", "Fault ports", "接收光功率", "RX power"],
            [None, None, None, None, "发射光功率", "TX power"],
            [None, None, "GNSS 信息", "GNSS info", "工作模式", "Operating mode"],
        ], ["A2:A4", "B2:B4", "C2:C3", "D2:D3"])
        result = preview_import(data, "functions.xlsx")
        self.assertEqual(result["summary"]["functions"], 1)
        self.assertEqual(result["summary"]["objects"], 2)
        self.assertEqual(result["summary"]["fields"], 3)
        function = result["catalog"]["functions"][0]
        self.assertEqual([item["name_zh"] for item in function["objects"][0]["fields"]], ["接收光功率", "发射光功率"])
        self.assertEqual(function["objects"][1]["fields"][0]["sources"], [{"sheet": "功能清单", "row": 4}])

    def test_new_question_resets_child_object(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "光功率", "Power"],
            ["查告警", "Query alarms", None, None, "告警名称", "Alarm name"],
            [None, None, "告警表", "Alarms", "状态", "Status"],
        ])
        result = preview_import(data, "functions.xlsx")
        functions = result["catalog"]["functions"]
        self.assertEqual(len(functions), 2)
        self.assertEqual([item["name_zh"] for item in functions[1]["objects"]], ["告警表"])
        self.assertEqual([item["name_zh"] for item in functions[1]["objects"][0]["fields"]], ["状态"])
        self.assertTrue(any("缺少对象" in warning for warning in result["warnings"]))

    def test_blank_language_not_copied_between_groups(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"],
            ["查告警", None, "告警", None, "名称", "Name"],
        ])
        result = preview_import(data, "functions.xlsx")
        second = result["catalog"]["functions"][1]
        self.assertEqual(second["question_en"], "")
        self.assertEqual(second["objects"][0]["name_en"], "")

    def test_inconsistent_language_merges_not_mixed(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"],
            ["查告警", None, "告警", None, "名称", "Name"],
        ], ["B2:B3", "D2:D3"])
        result = preview_import(data, "functions.xlsx")
        second = result["catalog"]["functions"][1]
        self.assertEqual(second["question_en"], "")
        self.assertEqual(second["objects"][0]["name_en"], "")
        self.assertTrue(any("合并范围不一致" in warning for warning in result["warnings"]))

    def test_english_only_rows_and_objects_without_fields(self):
        data = workbook_bytes([
            HEADERS,
            [None, "Analyze ticket", None, "Topology", None, None],
            [None, None, None, "Ports", None, "VSWR"],
        ])
        result = preview_import(data, "functions.xlsx")
        function = result["catalog"]["functions"][0]
        self.assertEqual(function["question_zh"], "")
        self.assertEqual(function["question_en"], "Analyze ticket")
        self.assertEqual(len(function["objects"]), 2)
        self.assertEqual(function["objects"][0]["fields"], [])

    def test_separator_resets_context_and_repeated_headers_are_skipped(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"],
            [None] * 6,
            [None, None, None, None, "不该继承的字段", "Unknown"],
            HEADERS,
            ["查告警", "Query alarms", "告警", "Alarms", "名称", "Name"],
        ])
        result = preview_import(data, "functions.xlsx")
        self.assertEqual(result["summary"]["fields"], 2)
        self.assertEqual(result["summary"]["skipped_rows"], 2)
        self.assertEqual(result["summary"]["blank_rows"], 1)

    def test_csv_utf8_and_forward_filled_hierarchy(self):
        data = (",".join(HEADERS) + "\n分析工单,Analyze ticket,端口,Port,接收光功率,RX\n,,,,发射光功率,TX\n").encode("utf-8-sig")
        result = preview_import(data, "functions.csv")
        self.assertEqual(result["summary"]["fields"], 2)
        self.assertEqual(result["sheets"], ["CSV"])
        self.assertEqual(result["header_row"], 1)

    def test_column_mapping_and_no_header(self):
        data = "接收功率,RX,端口,Port,分析工单,Analyze ticket\n".encode()
        mapping = {"question_zh": 5, "question_en": 6, "object_zh": 3, "object_en": 4, "field_zh": 1, "field_en": 2}
        result = preview_import(data, "functions.csv", columns=mapping, header_row=0)
        self.assertEqual(result["catalog"]["functions"][0]["question_zh"], "分析工单")
        self.assertEqual(result["summary"]["fields"], 1)

    def test_header_detection_and_auto_sheet(self):
        data = workbook_bytes([
            ["FMEMate 功能映射"],
            ["以下填写六列中英文信息"],
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"],
        ], ["A1:F1"], prefix_sheet=True)
        result = preview_import(data, "functions.xlsx")
        self.assertEqual(result["selected_sheet"], "功能清单")
        self.assertEqual(result["header_row"], 3)
        self.assertEqual(result["columns"], DEFAULT_COLUMNS)
        self.assertEqual(result["sheets"], ["填写说明", "功能清单"])

    def test_horizontal_semantic_merge_rejected(self):
        data = workbook_bytes([HEADERS, ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"]], ["C2:D2"])
        with self.assertRaisesRegex(ValueError, "横向合并"):
            preview_import(data, "functions.xlsx")

    def test_formula_and_error_rows_are_skipped(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"],
            [None, None, None, None, "=1+1", "TX"],
            [None, None, None, None, "不得继承", "Should not import"],
            ["查告警", "Query alarms", "告警", "Alarms", "#REF!", "Name"],
        ])
        result = preview_import(data, "functions.xlsx")
        self.assertEqual(result["summary"]["fields"], 1)
        self.assertTrue(any("公式或 Excel 错误值" in warning for warning in result["warnings"]))

    def test_deduplicates_repeated_rows_and_ids_are_stable(self):
        row = ["分析工单", "Analyze ticket", "端口", "Port", "接收功率", "RX"]
        data = workbook_bytes([HEADERS, row, row])
        result = preview_import(data, "first.xlsx")
        other = preview_import(data, "second.xlsx")
        self.assertEqual(result["summary"]["fields"], 1)
        field = result["catalog"]["functions"][0]["objects"][0]["fields"][0]
        self.assertEqual(len(field["sources"]), 2)
        self.assertEqual(result["catalog"]["functions"], other["catalog"]["functions"])

    def test_input_validation(self):
        with self.assertRaisesRegex(ValueError, "非空"):
            preview_import(b"", "empty.xlsx")
        with self.assertRaisesRegex(ValueError, "旧版"):
            preview_import(b"abc", "old.xls")
        with self.assertRaisesRegex(ValueError, "损坏"):
            preview_import(b"not a zip", "bad.xlsx")
        with self.assertRaisesRegex(ValueError, "不同"):
            preview_import(b"data", "x.csv", columns={key: 1 for key in DEFAULT_COLUMNS})
        with self.assertRaisesRegex(ValueError, "表头行"):
            preview_import(b"data", "x.csv", header_row=True)
        with self.assertRaisesRegex(ValueError, "不存在"):
            preview_import(b"data", "x.csv", sheet="missing")

    def test_merge_cannot_leak_across_parent_on_continuation_rows(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "光功率", "Power"],
            ["查告警", "Query alarms", None, None, "告警名称", "Alarm name"],
            [None, None, None, None, "告警状态", "Alarm status"],
        ], ["C2:C4", "D2:D4"])
        result = preview_import(data, "functions.xlsx")
        self.assertEqual(result["catalog"]["functions"][1]["objects"], [])
        self.assertEqual(result["summary"]["fields"], 1)

    def test_field_merge_cannot_leak_across_object(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "功率", "Power"],
            [None, None, "GNSS", "GNSS", None, None],
            [None, None, None, None, None, None],
        ], ["E2:E4", "F2:F4"])
        result = preview_import(data, "functions.xlsx")
        second = result["catalog"]["functions"][0]["objects"][1]
        self.assertEqual(second["fields"], [])

    def test_punctuation_is_preserved_in_identity(self):
        data = workbook_bytes([
            HEADERS,
            ["分析工单", "Analyze ticket", "端口", "Port", "A+B", "A+B"],
            [None, None, None, None, "AB", "AB"],
        ])
        result = preview_import(data, "functions.xlsx")
        self.assertEqual(result["summary"]["fields"], 2)

    def test_zip_expansion_limit(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"a" * (121 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "解压后过大"):
            preview_import(stream.getvalue(), "too_large.xlsx")


if __name__ == "__main__":
    unittest.main()
