"""Import the function -> object -> field hierarchy without losing Excel merges."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

MAX_UPLOAD = 20 * 1024 * 1024
MAX_EXPANDED = 120 * 1024 * 1024
MAX_ROWS = 100_000
MAX_SHEETS = 20
KEYS = ("question_zh", "question_en", "object_zh", "object_en", "field_zh", "field_en")
DEFAULT_COLUMNS = {key: index + 1 for index, key in enumerate(KEYS)}


def _norm(value: Any) -> str:
    return re.sub(r"[\s\W_]+", "", unicodedata.normalize("NFKC", str(value or "")).lower())


_HEADER_ALIASES = {
    "question_zh": ("触发问题", "触发问题中文", "中文触发问题", "问题中文", "中文问题", "问题", "预置问题", "预置问题中文", "功能中文", "questionzh", "triggerquestionzh", "chinesequestion"),
    "question_en": ("英文的触发问题", "触发问题英文", "英文触发问题", "英文问题", "问题英文", "预置问题英文", "功能英文", "questionen", "triggerquestionen", "englishquestion", "triggerquestionenglish"),
    "object_zh": ("对象", "对象中文", "中文对象", "对象名称", "对象名称中文", "数据对象", "表对象", "objectzh", "chineseobject"),
    "object_en": ("英文对象", "对象英文", "对象名称英文", "英文对象名称", "objecten", "englishobject", "objectenglish"),
    "field_zh": ("字段", "字段中文", "中文字段", "字段名称", "字段名称中文", "fieldzh", "chinesefield"),
    "field_en": ("字段英文", "英文字段", "字段名称英文", "英文字段名称", "fielden", "englishfield", "fieldenglish"),
}
_HEADER_LOOKUP = {_norm(alias): key for key, aliases in _HEADER_ALIASES.items() for alias in aliases}
# Generic English headers denote the English content column when accompanied by Chinese columns.
_HEADER_LOOKUP.update({"triggerquestion": "question_en", "question": "question_en", "object": "object_en", "objectname": "object_en", "field": "field_en", "fieldname": "field_en"})


def _header_mapping(values: list[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(values):
        key = _HEADER_LOOKUP.get(_norm(value))
        if key and key not in result:
            result[key] = index + 1
    return result


def _is_header(values: list[Any]) -> bool:
    return len(_header_mapping(values)) >= 3


def _text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def _label_key(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value).casefold()).strip()


def _identifier(prefix: str, *parts: str) -> str:
    payload = json.dumps([_label_key(part) for part in parts], ensure_ascii=False, separators=(",", ":"))
    return prefix + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def _source(target: dict, sheet: str, row: int) -> None:
    reference = {"sheet": sheet, "row": row}
    if reference not in target["sources"]:
        target["sources"].append(reference)


def _validate_columns(columns: dict | None) -> dict[str, int] | None:
    if columns is None:
        return None
    if not isinstance(columns, dict) or set(columns) != set(KEYS):
        raise ValueError("列映射必须包含 question_zh、question_en、object_zh、object_en、field_zh、field_en 六列。")
    result = {}
    for key, number in columns.items():
        if isinstance(number, bool) or not isinstance(number, int) or not 1 <= number <= 256:
            raise ValueError("列号必须是 1–256 的整数（A 列为 1）。")
        result[key] = number
    if len(set(result.values())) != len(KEYS):
        raise ValueError("六个字段必须分别选择不同的列。")
    return result


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV 后重试。")


def _inspect_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000 or sum(info.file_size for info in infos) > MAX_EXPANDED:
                raise ValueError("Excel 解压后过大，最多允许 120 MB / 5000 个内部文件。")
            for info in infos:
                if info.flag_bits & 1:
                    raise ValueError("不支持加密的 Excel，请先取消密码保护。")
    except zipfile.BadZipFile as exc:
        raise ValueError("Excel 文件损坏，或文件扩展名与实际格式不一致。") from exc


def _load_source(content: bytes, filename: str) -> tuple[list[str], dict[str, Any], Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        decoded = _decode_csv(content)
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = []
        try:
            for row in csv.reader(io.StringIO(decoded), dialect):
                if len(row) > 256:
                    raise ValueError("CSV 列数超过 256，请保留功能配置所需列。")
                rows.append(row)
                if len(rows) > MAX_ROWS:
                    raise ValueError("最多支持 100000 行，请拆分后导入。")
        except csv.Error as exc:
            raise ValueError("CSV 格式无法解析，请检查引号和分隔符。") from exc
        return ["CSV"], {"CSV": rows}, None
    if suffix not in (".xlsx", ".xlsm"):
        raise ValueError("支持 .xlsx、.xlsm 和 .csv；旧版 .xls 请先另存为 .xlsx。")
    _inspect_zip(content)
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False, keep_links=False, keep_vba=False)
    except ImportError as exc:
        raise ValueError("缺少 openpyxl，请先运行 pip install -r requirements.txt。") from exc
    except Exception as exc:
        raise ValueError("Excel 无法读取，请确认文件没有损坏或加密。") from exc
    if len(workbook.worksheets) > MAX_SHEETS:
        workbook.close()
        raise ValueError("最多支持 20 个工作表，请删除无关工作表后导入。")
    if sum(worksheet.max_row for worksheet in workbook.worksheets) > MAX_ROWS:
        workbook.close()
        raise ValueError("工作簿总行数超过 100000，请拆分文件或清除多余行格式。")
    if any(worksheet.max_column > 256 for worksheet in workbook.worksheets):
        workbook.close()
        raise ValueError("Excel 列数超过 256，请清除多余列格式后导入。")
    return workbook.sheetnames, {worksheet.title: worksheet for worksheet in workbook.worksheets}, workbook


def _first_rows(source: Any) -> list[list[Any]]:
    if isinstance(source, list):
        return source[:50]
    return [[cell.value for cell in row] for row in source.iter_rows(max_row=min(source.max_row, 50))]


def _select_sheet(sheets: list[str], sources: dict, requested: str | None) -> str:
    if requested is not None:
        if requested not in sheets:
            raise ValueError("所选工作表不存在，请重新选择。")
        return requested
    for name in sheets:
        if any(_is_header(row) for row in _first_rows(sources[name])):
            return name
    for name in sheets:
        if any(any(_text(value) for value in row) for row in _first_rows(sources[name])):
            return name
    raise ValueError("文件中没有可导入的数据。")


def _parse_rows(source: Any, name: str, columns: dict, start_row: int) -> tuple[list[dict], list[str]]:
    rows = []
    warnings = []
    merges: dict[tuple[int, int], tuple[Any, int, str]] = {}
    selected = set(columns.values())
    if not isinstance(source, list):
        for merged in source.merged_cells.ranges:
            if merged.min_row < start_row and merged.max_row < start_row:
                continue
            covered = selected.intersection(range(merged.min_col, merged.max_col + 1))
            if len(covered) > 1:
                raise ValueError(f"工作表 {name} 的 {merged} 横向合并跨越了多个语义列，请取消该横向合并后导入。")
            if merged.max_row > MAX_ROWS:
                raise ValueError("合并单元格超过 100000 行限制。")
            if covered:
                anchor = source.cell(merged.min_row, merged.min_col)
                for row in range(max(merged.min_row, start_row), merged.max_row + 1):
                    for col in covered:
                        merges[row, col] = (anchor.value, merged.min_row, anchor.data_type)
        total = source.max_row
    else:
        total = len(source)
    for row_number in range(start_row, total + 1):
        values = {}
        explicit = {}
        origins = {}
        invalid = False
        for key, column in columns.items():
            if isinstance(source, list):
                row = source[row_number - 1]
                value = row[column - 1] if len(row) >= column else None
                origin = row_number
                datatype = "f" if isinstance(value, str) and value.lstrip().startswith("=") else "s"
            else:
                cell = source.cell(row_number, column)
                value, origin, datatype = merges.get((row_number, column), (cell.value, row_number, cell.data_type))
            if datatype in ("f", "e"):
                invalid = True
            values[key] = _text(value)
            explicit[key] = bool(values[key]) and origin == row_number
            origins[key] = origin
        rows.append({"row": row_number, "values": values, "explicit": explicit, "origins": origins, "invalid": invalid})
    return rows, warnings


def _pair(values: dict, explicit: dict, prefix: str, current: tuple[str, str] | None, warnings: list[str], row_number: int) -> tuple[tuple[str, str] | None, bool]:
    keys = prefix + "_zh", prefix + "_en"
    candidate = (values[keys[0]], values[keys[1]])
    direct = (explicit[keys[0]], explicit[keys[1]])
    if not any(candidate):
        return current, False
    # A new Chinese label must never acquire an English label inherited from an older merge.
    # Equally, a new English-only group must not inherit an older Chinese label.
    changed_direct = [index for index in (0, 1) if direct[index] and (current is None or _label_key(candidate[index]) != _label_key(current[index]))]
    if changed_direct:
        candidate = tuple(candidate[index] if direct[index] or index in changed_direct else "" for index in (0, 1))
        if any(values[keys[index]] and not candidate[index] for index in (0, 1)):
            warnings.append(f"第 {row_number} 行的中英文合并范围不一致，已忽略来自上一分组的名称，请确认预览。")
        return candidate, current is None or tuple(map(_label_key, candidate)) != tuple(map(_label_key, current))
    if current is None:
        return candidate, True
    return current, False


def preview_import(content: bytes, filename: str, sheet: str | None = None, columns: dict | None = None, header_row: int | None = None) -> dict:
    """Return an import preview; callers explicitly confirm before persisting its catalog.

    Column indexes are 1-based. ``header_row=None`` detects a header in the first
    50 rows, and ``header_row=0`` means the data starts in row 1.
    """
    if not isinstance(content, bytes) or not content:
        raise ValueError("请选择非空的 Excel 或 CSV 文件。")
    if len(content) > MAX_UPLOAD:
        raise ValueError("文件超过 20 MB，请拆分后导入。")
    if not isinstance(filename, str) or not filename.strip():
        raise ValueError("文件名不能为空。")
    if header_row is not None and (isinstance(header_row, bool) or not isinstance(header_row, int) or not 0 <= header_row <= MAX_ROWS):
        raise ValueError("表头行必须为 0–100000 的整数，0 表示无表头。")
    mapped_columns = _validate_columns(columns)
    sheets, sources, workbook = _load_source(content, filename)
    try:
        selected_sheet = _select_sheet(sheets, sources, sheet)
        source = sources[selected_sheet]
        initial_rows = _first_rows(source)
        detected_header = next((index + 1 for index, row in enumerate(initial_rows) if _is_header(row)), 0)
        selected_header = detected_header if header_row is None else header_row
        warnings: list[str] = []
        if mapped_columns is None:
            if selected_header and selected_header <= len(initial_rows):
                detected = _header_mapping(initial_rows[selected_header - 1])
                if set(detected) == set(KEYS):
                    mapped_columns = detected
                else:
                    mapped_columns = DEFAULT_COLUMNS.copy()
                    warnings.append("表头未完整识别，暂按前六列：触发问题中文/英文、对象中文/英文、字段中文/英文。请确认列映射。")
            else:
                mapped_columns = DEFAULT_COLUMNS.copy()
                warnings.append("未识别到表头，暂按前六列导入。请确认表头行与列映射。")
        if selected_header == 0:
            warnings.append("当前按无表头文件导入，第一行也会作为数据。")
        if len(sheets) > 1:
            warnings.append(f"当前仅导入工作表「{selected_sheet}」；其余 {len(sheets) - 1} 个工作表未导入，可在上方切换。")
        row_records, row_warnings = _parse_rows(source, selected_sheet, mapped_columns, selected_header + 1)
        warnings.extend(row_warnings)
        functions: list[dict] = []
        function_map: dict[str, dict] = {}
        object_map: dict[str, dict] = {}
        field_map: dict[str, dict] = {}
        current_question: tuple[str, str] | None = None
        current_object: tuple[str, str] | None = None
        question_start_row = object_start_row = 0
        imported_rows = 0
        skipped_rows = 0
        blank_rows = 0
        for record in row_records:
            number, values, explicit = record["row"], record["values"], record["explicit"]
            if record["invalid"]:
                warnings.append(f"第 {number} 行含公式或 Excel 错误值，已跳过该行并重置分组，请改为文本值后导入。")
                current_question = current_object = None
                skipped_rows += 1
                continue
            if not any(values.values()):
                current_question = current_object = None
                blank_rows += 1
                continue
            if _is_header(list(values.values())):
                current_question = current_object = None
                skipped_rows += 1
                warnings.append(f"第 {number} 行为重复表头，已跳过并重置分组。")
                continue
            next_question, new_question = _pair(values, explicit, "question", current_question, warnings, number)
            current_question = next_question
            if new_question:
                current_object = None
                question_start_row = number
            if current_question is None:
                warnings.append(f"第 {number} 行缺少触发问题，已跳过；空白分隔行后需重新填写问题。")
                skipped_rows += 1
                continue
            function_id = _identifier("fn_", *current_question)
            if function_id not in function_map:
                function = {"id": function_id, "question_zh": current_question[0], "question_en": current_question[1], "objects": [], "sources": []}
                function_map[function_id] = function
                functions.append(function)
            function = function_map[function_id]
            _source(function, selected_sheet, number)
            # A merge that began under an earlier question must never reappear as
            # a child of the current question, including on later continuation rows.
            object_values, object_explicit = dict(values), dict(explicit)
            for key in ("object_zh", "object_en"):
                if record["origins"][key] < question_start_row:
                    object_values[key] = ""
                    object_explicit[key] = False
            current_object, new_object = _pair(object_values, object_explicit, "object", current_object, warnings, number)
            if new_object:
                object_start_row = number
            if current_object is None:
                if values["field_zh"] or values["field_en"]:
                    warnings.append(f"第 {number} 行有字段但缺少对象，字段未导入，请补齐对象。")
                    skipped_rows += 1
                else:
                    imported_rows += 1
                continue
            object_id = _identifier("obj_", *current_question, *current_object)
            if object_id not in object_map:
                obj = {"id": object_id, "name_zh": current_object[0], "name_en": current_object[1], "fields": [], "sources": []}
                object_map[object_id] = obj
                function["objects"].append(obj)
            obj = object_map[object_id]
            _source(obj, selected_sheet, number)
            field_values, field_explicit = dict(values), dict(explicit)
            for key in ("field_zh", "field_en"):
                if record["origins"][key] < object_start_row:
                    field_values[key] = ""
                    field_explicit[key] = False
            field_names, _ = _pair(field_values, field_explicit, "field", None, warnings, number)
            if field_names:
                field_id = _identifier("fld_", *current_question, *current_object, *field_names)
                if field_id not in field_map:
                    field = {"id": field_id, "name_zh": field_names[0], "name_en": field_names[1], "sources": []}
                    field_map[field_id] = field
                    obj["fields"].append(field)
                _source(field_map[field_id], selected_sheet, number)
            imported_rows += 1
        if not functions:
            raise ValueError("未找到有效的触发问题。请检查工作表、表头行和六列映射。")
        if not object_map:
            warnings.append("本次仅导入了问题，没有有效对象。字段推荐需先补齐对象。")
        metadata = {"filename": Path(filename).name, "sheet": selected_sheet, "header_row": selected_header, "columns": mapped_columns, "source_sha256": hashlib.sha256(content).hexdigest()}
        return {
            "catalog": {"functions": functions, "metadata": metadata},
            "summary": {"functions": len(functions), "objects": len(object_map), "fields": len(field_map), "rows": imported_rows, "skipped_rows": skipped_rows, "blank_rows": blank_rows},
            "warnings": list(dict.fromkeys(warnings))[:200],
            "sheets": sheets,
            "selected_sheet": selected_sheet,
            "columns": mapped_columns,
            "header_row": selected_header,
        }
    finally:
        if workbook is not None:
            workbook.close()
